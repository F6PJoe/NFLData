#!/usr/bin/env python3
"""
One-off exploration: build team-level EPA/play and Yards/play (offense
generated, defense allowed) from nflverse's free, public play-by-play data,
and lay it out next to the current Live sheet's Def DAVE / Opp Off DAVE
columns so Joe can see how it compares before deciding whether to replace
FTN DAVE with this.

Data: nflverse-data's "pbp" GitHub release (play_by_play_<year>.csv.gz) --
no login, no scraping, just a public data file, updated same-day
throughout the season. `epa` is nflfastR's own precomputed per-play value
(see the EPA formula discussion in chat) -- this script doesn't reimplement
the expected-points model, it aggregates the already-computed numbers.

Methodology (ours, not borrowed):
  - Only "pass"/"run" plays count (scrimmage plays -- excludes kickoffs,
    punts, FGs, kneels/spikes, which have a different EPA scale and aren't
    meaningful "defense allowed" plays in the same sense).
  - Two versions of each stat: ALL plays, and a garbage-time-excluded
    version (drops plays where the offense's win probability is outside
    5%-95% -- a standard convention for "the game was still competitive").
  - Rank convention matches the rest of the sheet: for DEFENSE, sorting
    worst-to-best (highest EPA/yards allowed first) puts the worst defense
    at rank 1 and the best at rank <n> -- same direction as the existing
    "Def DAVE" column. For OFFENSE, sorting best-to-worst (highest EPA/
    yards generated first) puts the best offense at rank 1 -- same
    direction as "Opp Off DAVE" (used as-is for the opponent, no
    inversion, since a low-number rank there already means a tough
    opponent offense and a high number means a weak one worth exploiting).

Usage:
    python explore_nflverse_epa.py
    -> nflverse_epa_exploration.xlsx

Requires: pandas, openpyxl, requests, google-api-python-client, google-auth
"""

import os

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from team_names import normalize

PBP_URL = "https://github.com/nflverse/nflverse-data/releases/download/pbp/play_by_play_{year}.csv.gz"
YEAR = 2026

SHEET_ID = "1lTRoatl-YQHlv78YeisG7eFyz2xqaConGUoPo4medpQ"
SERVICE_ACCT = os.path.join(os.path.dirname(__file__), "..", "triple-baton-456523-e4-b9ec3cbd6e3d.json")
TAB = "Live"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="333333")


def fetch_pbp():
    cache_path = f"nflverse_data/play_by_play_{YEAR}.csv.gz"
    if not os.path.exists(cache_path):
        os.makedirs("nflverse_data", exist_ok=True)
        resp = requests.get(PBP_URL.format(year=YEAR), timeout=60)
        resp.raise_for_status()
        with open(cache_path, "wb") as f:
            f.write(resp.content)
    df = pd.read_csv(cache_path, compression="gzip", low_memory=False)
    return df[df["play_type"].isin(["pass", "run"])].copy()


def team_rollups(plays):
    """Returns (defense_df, offense_df), each indexed by team abbr."""
    competitive = plays[(plays["wp"] >= 0.05) & (plays["wp"] <= 0.95)]

    def rollup(df, group_col, prefix):
        g_all = df.groupby(group_col).agg(
            **{f"{prefix}_epa_all": ("epa", "mean"), f"{prefix}_ypp_all": ("yards_gained", "mean")}
        )
        comp = competitive.groupby(group_col).agg(
            **{f"{prefix}_epa_comp": ("epa", "mean"), f"{prefix}_ypp_comp": ("yards_gained", "mean")}
        )
        plays_n = df.groupby(group_col).size().rename(f"{prefix}_plays")
        return g_all.join(comp).join(plays_n)

    defense = rollup(plays, "defteam", "def")
    offense = rollup(plays, "posteam", "off")
    return defense, offense


def add_ranks(df, epa_col, ypp_col):
    """Ranks 1..n by descending raw value, which lands on the sheet's own
    conventions for both sides (confirmed with Joe):

      DEFENSE: value is EPA/yards ALLOWED, so descending = worst defense
               first -> rank 1 = worst D, rank 32 = best D. Matches the
               "Def DAVE" column (33 - FTN's 1=best rank).
      OFFENSE: value is EPA/yards GENERATED, so descending = best offense
               first -> rank 1 = best O, rank 32 = worst O. Matches the
               "Opp Off DAVE" column (FTN's Off Rank used as-is).
    """
    df = df.sort_values(epa_col, ascending=False)
    df["EPA Rank"] = range(1, len(df) + 1)
    df = df.sort_values(ypp_col, ascending=False)
    df["Yards Rank"] = range(1, len(df) + 1)
    return df


def fetch_live_sheet():
    from google.oauth2.service_account import Credentials
    from googleapiclient.discovery import build

    creds = Credentials.from_service_account_file(
        SERVICE_ACCT, scopes=["https://www.googleapis.com/auth/spreadsheets"]
    )
    service = build("sheets", "v4", credentials=creds, cache_discovery=False)
    resp = service.spreadsheets().values().get(
        spreadsheetId=SHEET_ID, range=f"{TAB}!B2:I"
    ).execute()
    rows = resp.get("values", [])
    out = []
    for row in rows:
        if len(row) < 8:
            continue
        out.append({
            "Team": row[0], "Opp": row[3],
            "Def DAVE (current)": row[6], "Opp Off DAVE (current)": row[7],
        })
    return pd.DataFrame(out)


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def autosize(ws):
    for col_cells in ws.columns:
        length = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 2, 28)


def main():
    print("Fetching nflverse play-by-play...")
    plays = fetch_pbp()
    print(f"{len(plays)} scrimmage plays loaded (weeks {sorted(plays['week'].unique())}).")

    defense, offense = team_rollups(plays)
    defense = add_ranks(defense, "def_epa_all", "def_ypp_all")
    offense = add_ranks(offense, "off_epa_all", "off_ypp_all")

    defense = defense.reset_index().rename(columns={"defteam": "Team"})
    offense = offense.reset_index().rename(columns={"posteam": "Team"})
    defense["Team"] = defense["Team"].apply(normalize)
    offense["Team"] = offense["Team"].apply(normalize)

    print("Fetching current Live sheet for comparison...")
    live = fetch_live_sheet()
    live["Team_norm"] = live["Team"].apply(normalize)
    live["Opp_norm"] = live["Opp"].apply(normalize)

    comparison = live.merge(
        defense[["Team", "def_epa_all", "Yards Rank", "EPA Rank"]].rename(
            columns={"Yards Rank": "New Yards Rank", "EPA Rank": "New EPA Rank", "def_epa_all": "EPA/Play Allowed"}
        ),
        left_on="Team_norm", right_on="Team", how="left", suffixes=("", "_def"),
    )
    comparison = comparison.merge(
        offense[["Team", "off_epa_all", "EPA Rank"]].rename(
            columns={"EPA Rank": "New Opp Off EPA Rank", "off_epa_all": "Opp EPA/Play"}
        ),
        left_on="Opp_norm", right_on="Team", how="left", suffixes=("", "_opp"),
    )
    comparison = comparison[[
        "Team", "Opp", "Def DAVE (current)", "New EPA Rank", "New Yards Rank",
        "Opp Off DAVE (current)", "New Opp Off EPA Rank",
    ]]

    wb = Workbook()
    wb.remove(wb.active)

    ws1 = wb.create_sheet("Defense EPA Model")
    def_cols = ["Team", "def_plays", "def_epa_all", "def_ypp_all", "def_epa_comp", "def_ypp_comp",
                "EPA Rank", "Yards Rank"]
    def_headers = ["Team", "Plays", "EPA/Play Allowed", "Yards/Play Allowed",
                   "EPA/Play Allowed (no garbage time)", "Yards/Play Allowed (no garbage time)",
                   "EPA Rank (1=worst D, 32=best D)", "Yards Rank (1=worst D, 32=best D)"]
    ws1.append(def_headers)
    for _, r in defense.sort_values("EPA Rank", ascending=False).iterrows():
        ws1.append([r[c] for c in def_cols])
    style_header(ws1)
    autosize(ws1)

    ws2 = wb.create_sheet("Offense EPA Model")
    off_cols = ["Team", "off_plays", "off_epa_all", "off_ypp_all", "off_epa_comp", "off_ypp_comp",
                "EPA Rank", "Yards Rank"]
    off_headers = ["Team", "Plays", "EPA/Play", "Yards/Play",
                   "EPA/Play (no garbage time)", "Yards/Play (no garbage time)",
                   "EPA Rank (1=best O, 32=worst O)", "Yards Rank (1=best O, 32=worst O)"]
    ws2.append(off_headers)
    for _, r in offense.sort_values("EPA Rank").iterrows():
        ws2.append([r[c] for c in off_cols])
    style_header(ws2)
    autosize(ws2)

    ws3 = wb.create_sheet("Vs Current DVOA")
    ws3.append(list(comparison.columns))
    for _, r in comparison.sort_values("New EPA Rank", ascending=False, na_position="last").iterrows():
        ws3.append(list(r.values))
    style_header(ws3)
    autosize(ws3)

    # Context block -- both rank sets are now on the SAME scale (32=best D),
    # so this correlation is meaningful. Low agreement here is expected
    # right now for reasons that have nothing to do with either metric
    # being wrong; spelled out so the number isn't read as alarming.
    games = plays.groupby("defteam")["game_id"].nunique()
    one_game = int((games == 1).sum())
    dave_vals = pd.to_numeric(comparison["Def DAVE (current)"], errors="coerce")
    new_vals = pd.to_numeric(comparison["New EPA Rank"], errors="coerce")
    r = dave_vals.corr(new_vals)

    ws3.append([])
    for line in [
        f"Rank correlation, current DAVE vs new EPA model: {r:.3f}",
        f"Sample: {one_game} of 32 teams have played only ONE game so far.",
        "FTN's own page: 'DAVE defense and special teams are 98% projection "
        "and 2% actual performance' this early -- so DAVE is still mostly "
        "preseason projection, while this model is 100% actual results.",
        "Low agreement right now is expected and is not evidence either "
        "metric is broken. Re-check around Week 6-8 for a fair comparison.",
    ]:
        ws3.append([line])
    for row in ws3.iter_rows(min_row=ws3.max_row - 3, max_row=ws3.max_row):
        row[0].font = Font(italic=True)

    out_path = "nflverse_epa_exploration.xlsx"
    wb.save(out_path)
    print(f"Wrote {out_path} ({len(defense)} teams).")


if __name__ == "__main__":
    main()
