#!/usr/bin/env python3
"""Sanity-check these projections against the ff_draft_proj consensus.

Consensus is a YARDSTICK, never an input. Nothing here feeds back into the
model - it reads consensus_*.csv only to report divergence, so you can tell
"I disagree with the market about this player" apart from "something is
broken." If this script ever starts writing back into the projection inputs,
it has defeated the point of building an independent model.

Three levels of check, cheapest signal first:

  1. LEAGUE TOTALS   are the aggregate targets/carries/yards/TDs in the right
                     universe at all? Catches a broken denominator instantly.
  2. RANK AGREEMENT  Spearman correlation vs consensus per position. High
                     agreement with a few big outliers is what you want.
                     Low correlation means something structural is wrong.
  3. PER PLAYER      biggest over- and under-projections vs consensus.

Reads the workbook's cached values if Excel has saved it; otherwise computes
from the seeded shares so you can check a starting point before touching it.

Usage:
    python compare_to_consensus.py
    python compare_to_consensus.py --top 25
"""

import argparse
import csv
import os
import sys
from collections import defaultdict

import names
from teams import clean_team

HERE = os.path.dirname(os.path.abspath(__file__))
CONSENSUS = os.path.join(HERE, os.pardir, "ff_draft_proj")
WORKBOOK = os.path.join(HERE, "my_projections.xlsx")

POSITIONS = ("QB", "RB", "WR", "TE")
CONSENSUS_POINTS = {           # the half-PPR column name varies by position
    "QB": "Fantasy Points",
    "RB": "Fantasy Points (Half-PPR)",
    "WR": "Fantasy Points (Half)",
    "TE": "Fantasy Points (Half-PPR)",
}
# Rough NFL reality, per season across all 32 teams. Used as an absolute
# floor/ceiling check that doesn't depend on consensus at all.
LEAGUE_TRUTH = {
    "targets": (16_500, 17_500),
    "rush_att": (13_000, 14_500),
    "rec_yds": (118_000, 130_000),
    "rush_yds": (58_000, 66_000),
    "rec_td": (750, 900),
    "rush_td": (400, 520),
}


def f(v, d=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return d


def half_ppr(rec_yds, rec_td, rec, rush_yds, rush_td,
             pass_yds=0.0, pass_td=0.0, pass_int=0.0):
    return (rec_yds * 0.1 + rec_td * 6 + rec * 0.5
            + rush_yds * 0.1 + rush_td * 6
            + pass_yds * 0.04 + pass_td * 4 - pass_int * 2)


def from_seeds():
    """Recompute the workbook's arithmetic in Python from the seed data."""
    shares = {r["player"] + "|" + r["team"]: r for r in
              csv.DictReader(open(os.path.join(HERE, "player_shares.csv"),
                                  encoding="utf-8"))}
    totals = {r["team"]: r for r in
              csv.DictReader(open(os.path.join(HERE, "team_totals_2026.csv"),
                                  encoding="utf-8"))}
    target_rate = 0.954
    out = []
    for p in shares.values():
        if p["status"] == "OUT":
            continue
        t = totals.get(p["team"])
        if not t:
            continue
        tgt_share, rush_share = f(p["tgt_share_w"]), f(p["rush_share_w"])
        # Two kinds of "no starting value", both of which need a human, and
        # neither of which is a disagreement with consensus:
        #   changed_team - prior share belonged to a different offense
        #   rookie       - no NFL history at all to seed from
        changed = p.get("changed_team") == "Y"
        rookie = not any(f(p.get(f"{k}_w")) for k in
                         ("tgt_share", "rush_share", "pass_share"))
        unfilled = changed or rookie
        reason = "changed team" if changed else "rookie/no snaps" if rookie else ""
        if unfilled:
            tgt_share = rush_share = 0.0
        team_tgt = round(f(t["pass_att"]) * target_rate)
        targets = tgt_share * team_tgt
        rec = targets * f(p["catch_pct_w"])
        rec_yds = rec * f(p["ypr_w"])
        rec_td = tgt_share * f(p["rz_tgt_ratio"], 1.0) * f(t["pass_td"])
        rush_att = rush_share * f(t["rush_att"])
        rush_yds = rush_att * f(p["ypc_w"])
        rush_td = rush_share * f(p["rz_rush_ratio"], 1.0) * f(t["rush_td"])
        pass_share = 0.0 if unfilled else f(p.get("pass_share_w"))
        pass_yds = pass_share * f(t["pass_yds"])
        pass_td = pass_share * f(t["pass_td"])
        pass_int = pass_share * f(t["interceptions"])
        out.append({
            "player": p["player"], "team": p["team"],
            "pos": p["pos"] if p["pos"] != "FB" else "RB",
            "targets": targets, "rec": rec, "rec_yds": rec_yds,
            "rec_td": rec_td, "rush_att": rush_att, "rush_yds": rush_yds,
            "rush_td": rush_td, "unfilled": unfilled, "reason": reason,
            "pass_yds": pass_yds, "pass_td": pass_td, "pass_int": pass_int,
            "pts": half_ppr(rec_yds, rec_td, rec, rush_yds, rush_td,
                            pass_yds, pass_td, pass_int),
        })
    return out, "seeded shares (workbook not yet opened in Excel)"


def from_workbook():
    """Read cached values out of the workbook, if Excel has computed them."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None
    if not os.path.exists(WORKBOOK):
        return None
    wb = load_workbook(WORKBOOK, data_only=True)
    rows = []
    for pos in POSITIONS:
        if pos not in wb.sheetnames:
            return None
        ws = wb[pos]
        for r in range(3, ws.max_row + 1):
            name = ws.cell(r, 1).value
            if not name or not isinstance(name, str):
                continue
            vals = [ws.cell(r, c).value for c in range(4, 12)]
            if any(v is None for v in vals):
                return None            # formulas never evaluated - bail out
            tgt, rec, ry, rtd, ra, rushy, rushtd, pts = [f(v) for v in vals]
            rows.append({"player": name, "team": ws.cell(r, 2).value, "pos": pos,
                         "targets": tgt, "rec": rec, "rec_yds": ry,
                         "rec_td": rtd, "rush_att": ra, "rush_yds": rushy,
                         "rush_td": rushtd, "pts": pts, "unfilled": False})
    return (rows, "my_projections.xlsx (values as last saved by Excel)") if rows else None


def load_consensus():
    out = {}
    for pos in POSITIONS:
        path = os.path.join(CONSENSUS, f"consensus_{pos.lower()}.csv")
        if not os.path.exists(path):
            sys.exit(f"Missing {path} - run ff_draft_proj/build_consensus.py")
        with open(path, newline="", encoding="utf-8") as fh:
            for row in csv.DictReader(fh):
                name = row[pos]
                out[names.key(name, clean_team(row["Team"]))] = {
                    "player": name, "pos": pos,
                    "team": clean_team(row["Team"]),
                    "targets": f(row.get("Targets")),
                    "rec": f(row.get("Rec")),
                    "rec_yds": f(row.get("Rec Yds")),
                    "rec_td": f(row.get("Rec TD")),
                    "rush_att": f(row.get("Rush Att")),
                    "rush_yds": f(row.get("Rush Yds")),
                    "rush_td": f(row.get("Rush TD")),
                    "pass_yds": f(row.get("Pass Yds")),
                    "pass_td": f(row.get("Pass TD")),
                    "pass_int": f(row.get("Pass Int")),
                    "pts": f(row.get(CONSENSUS_POINTS[pos])),
                }
    return out


def spearman(pairs):
    """Rank correlation. Ties get average ranks."""
    if len(pairs) < 3:
        return 0.0

    def rank(vals):
        order = sorted(range(len(vals)), key=lambda i: -vals[i])
        r = [0.0] * len(vals)
        i = 0
        while i < len(order):
            j = i
            while j + 1 < len(order) and vals[order[j + 1]] == vals[order[i]]:
                j += 1
            avg = (i + j) / 2 + 1
            for k in range(i, j + 1):
                r[order[k]] = avg
            i = j + 1
        return r

    rx, ry = rank([p[0] for p in pairs]), rank([p[1] for p in pairs])
    n = len(pairs)
    mx, my = sum(rx) / n, sum(ry) / n
    sxy = sum((rx[i] - mx) * (ry[i] - my) for i in range(n))
    sxx = sum((v - mx) ** 2 for v in rx)
    syy = sum((v - my) ** 2 for v in ry)
    return sxy / (sxx * syy) ** 0.5 if sxx and syy else 0.0


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--top", type=int, default=15)
    args = ap.parse_args()

    loaded = from_workbook()
    mine, source = loaded if loaded else from_seeds()
    cons = load_consensus()

    print(f"Mine      : {source}")
    print(f"Consensus : ff_draft_proj/consensus_*.csv ({len(cons)} players)\n")

    # ── 1. league totals ───────────────────────────────────────────────────
    print("=" * 74)
    print("1. LEAGUE TOTALS - is this the right universe at all?")
    print("=" * 74)
    print(f"  {'':<12}{'mine':>12}{'consensus':>12}{'NFL actual':>22}")
    print(f"  {'-'*12}{'-'*12}{'-'*12}{'-'*22}")
    for key, (lo, hi) in LEAGUE_TRUTH.items():
        m = sum(p.get(key, 0.0) for p in mine)
        c = sum(p.get(key, 0.0) for p in cons.values())
        flag = "" if lo <= m <= hi else "   <-- low" if m < lo else "   <-- HIGH"
        print(f"  {key:<12}{m:>12,.0f}{c:>12,.0f}{lo:>11,.0f}-{hi:<10,.0f}{flag}")
    print("\n  Read this in context:")
    print("  - Unallocated share makes MINE read low until every team footer is")
    print("    at 100%. Seeded, teams average ~79% of targets, so expect roughly")
    print("    a 20% shortfall here. It closes as you fill teams in.")
    print("  - Consensus reads high on volume because it double-counts nothing")
    print("    but covers only draftable players with generous individual lines.")
    print("  - What matters: after balancing, mine should land INSIDE the range.")

    # ── 2. rank agreement ──────────────────────────────────────────────────
    print("\n" + "=" * 74)
    print("2. RANK AGREEMENT vs consensus (Spearman)")
    print("=" * 74)
    by_pos = defaultdict(list)
    unfilled = defaultdict(int)
    todo = []
    for p in mine:
        key = names.key(p["player"], p["team"])
        c = cons.get(key)
        if not c:
            continue
        if p.get("unfilled"):
            unfilled[p["pos"]] += 1
            todo.append((c["pts"], p, c))
            continue
        by_pos[p["pos"]].append((p, c))
    for pos in POSITIONS:
        pairs = [(p["pts"], c["pts"]) for p, c in by_pos[pos]]
        if not pairs:
            continue
        rho = spearman(pairs)
        verdict = ("strong" if rho >= 0.85 else "ok" if rho >= 0.7 else
                   "WEAK - investigate")
        extra = f", {unfilled[pos]} unfilled skipped" if unfilled[pos] else ""
        print(f"  {pos:<4} rho = {rho:>5.3f}   n = {len(pairs):<4} {verdict}{extra}")
    print("\n  0.85+ means you're ranking players the same way with different")
    print("  numbers - exactly what you want. Below 0.70 means something")
    print("  structural is off, not a difference of opinion.")

    # ── 3. per player ──────────────────────────────────────────────────────
    print("\n" + "=" * 74)
    print(f"3. BIGGEST DIVERGENCES (half PPR points)")
    print("=" * 74)
    diffs = []
    for pos in POSITIONS:
        for p, c in by_pos[pos]:
            if c["pts"] < 40:
                continue          # ignore deep bench noise
            diffs.append((p["pts"] - c["pts"], p, c))
    diffs.sort()
    for label, rows in (("MINE MUCH HIGHER than consensus", diffs[::-1][:args.top]),
                        ("MINE MUCH LOWER than consensus", diffs[:args.top])):
        print(f"\n  {label}:")
        print(f"    {'player':<24}{'tm':<5}{'pos':<5}{'mine':>8}{'cons':>8}{'diff':>8}")
        for d, p, c in rows:
            print(f"    {p['player'][:23]:<24}{p['team']:<5}{p['pos']:<5}"
                  f"{p['pts']:>8.0f}{c['pts']:>8.0f}{d:>+8.0f}")

    if todo:
        print("\n" + "=" * 74)
        print(f"4. NOT YET FILLED IN ({len(todo)}) - your actual to-do list")
        print("=" * 74)
        print("  These project 0 because there's nothing to seed from, not")
        print("  because the model disagrees. Excluded from the checks above.")
        print("  Ranked by what consensus thinks they're worth - work top down.\n")
        print(f"    {'player':<24}{'tm':<5}{'pos':<5}{'cons':>7}   why")
        todo.sort(key=lambda x: -x[0])
        for pts, p, c in todo[:args.top * 2]:
            print(f"    {p['player'][:23]:<24}{p['team']:<5}{p['pos']:<5}"
                  f"{pts:>7.0f}   {p.get('reason', '')}")
        by_reason = defaultdict(int)
        for _, p, _ in todo:
            by_reason[p.get("reason", "?")] += 1
        print("\n  " + ", ".join(f"{v} {k}" for k, v in sorted(by_reason.items())))


if __name__ == "__main__":
    sys.exit(main())
