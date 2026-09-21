#!/usr/bin/env python3
"""Read my_projections_v2.xlsx -> QB/RB/WR/TE ranking tabs, all 3 scoring formats.

v1 (my_projections.xlsx) already has live QB/RB/WR/TE tabs - cross-sheet
formulas that update as you type, since its output IS a normalized share of a
team total. v2 doesn't have that: it's raw typed numbers with no formula link
between team tabs, and v2's own workbook is already carrying 10k+ formulas and
960 conditional-format rules - adding 4 more LIVE cross-sheet tabs would only
add to the recalculation load that's already causing it to freeze.

So this is a script, not a tab: it reads the CURRENT values out of
my_projections_v2.xlsx (read-only - works even while the file is open in
Excel, no need to close it) and writes a small separate STATIC workbook.
Nothing here is a formula, so it costs nothing to keep around; re-run it
whenever you want the rankings refreshed.

Scoring reuses ff_draft_proj/scoring.py directly, not reimplemented - the same
rule the rest of this project follows, so there's exactly one scoring
definition in the repo and this can never quietly disagree with it.

Usage:
    python build_rankings.py
    python build_rankings.py --in my_projections.xlsx   # read v1 instead
"""

import argparse
import os
import sys
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                os.pardir, "ff_draft_proj"))
import scoring  # noqa: E402  (ff_draft_proj/scoring.py - the one true scoring rules)

HERE = os.path.dirname(os.path.abspath(__file__))

import build_workbook_v2 as wb2  # column layout (C, R_FIRST, R_LAST, TEAMS)

HDR = PatternFill("solid", fgColor="2F4F6F")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM1 = "0.0"

# Per position: (v2 source column, output column) pairs to pull from the
# workbook - explicit pairs, not same-name-assumed, because "Cmp"/"Int" in v2
# become "Pass Comp"/"Pass Int" in the output and a same-name lookup silently
# returns nothing for exactly those two columns (caught by inspecting the
# actual output, not by reasoning about it - worth checking real cells).
# Output layout matches ff_draft_proj/consensus_<pos>.csv exactly, so this
# reads the same as every other projection file in the repo.
POSITIONS = {
    "QB": {
        "stats": [("Pass Att", "Pass Att"), ("Cmp", "Pass Comp"),
                 ("Pass Yds", "Pass Yds"), ("Pass TD", "Pass TD"),
                 ("Int", "Pass Int"), ("Rush Att", "Rush Att"),
                 ("Rush Yds", "Rush Yds"), ("Rush TD", "Rush TD")],
        "out": ["Player", "Team", "Pass Att", "Pass Comp", "Pass Yds",
               "Pass TD", "Pass Int", "Rush Att", "Rush Yds", "Rush TD",
               "Fantasy Points"],
        "sort": "Fantasy Points",
    },
    "RB": {
        "stats": [("Rush Att", "Rush Att"), ("Rush Yds", "Rush Yds"),
                 ("Rush TD", "Rush TD"), ("Targets", "Targets"),
                 ("Rec", "Rec"), ("Rec Yds", "Rec Yds"), ("Rec TD", "Rec TD")],
        "out": ["Player", "Team", "Rush Att", "Rush Yds", "Rush TD",
               "Targets", "Rec", "Rec Yds", "Rec TD",
               "Fantasy Points (Half-PPR)", "Fantasy Points (PPR)",
               "Fantasy Points (STD)"],
        "sort": "Fantasy Points (Half-PPR)",
    },
    "WR": {
        "stats": [("Targets", "Targets"), ("Rec", "Rec"),
                 ("Rec Yds", "Rec Yds"), ("Rec TD", "Rec TD"),
                 ("Rush Att", "Rush Att"), ("Rush Yds", "Rush Yds"),
                 ("Rush TD", "Rush TD")],
        "out": ["Player", "Team", "Targets", "Rec", "Rec Yds", "Rec TD",
               "Rush Att", "Rush Yds", "Rush TD",
               "Fantasy Points (Half)", "Fantasy Points (PPR)",
               "Fantasy Points (STD)"],
        "sort": "Fantasy Points (Half)",
    },
    "TE": {
        "stats": [("Targets", "Targets"), ("Rec", "Rec"),
                 ("Rec Yds", "Rec Yds"), ("Rec TD", "Rec TD"),
                 ("Rush Att", "Rush Att"), ("Rush Yds", "Rush Yds"),
                 ("Rush TD", "Rush TD")],
        "out": ["Player", "Team", "Targets", "Rec", "Rec Yds", "Rec TD",
               "Rush Att", "Rush Yds", "Rush TD",
               "Fantasy Points (Half-PPR)", "Fantasy Points (PPR)",
               "TE Premium", "Fantasy Points (STD)"],
        "sort": "Fantasy Points (Half-PPR)",
    },
}


def read_players(path):
    """Every player row across all 32 team tabs, as {stat name: value}."""
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except FileNotFoundError:
        sys.exit(f"Missing {path}. Build it first.")
    players = defaultdict(list)  # pos -> list of row dicts
    for team in wb2.TEAMS:
        if team not in wb.sheetnames:
            continue
        ws = wb[team]
        for row in ws.iter_rows(min_row=wb2.R_FIRST, max_row=wb2.R_LAST):
            cell = {c.column: c.value for c in row}
            name = cell.get(wb2.C["Player"])
            pos = cell.get(wb2.C["Pos"])
            if not name or not isinstance(name, str):
                continue
            if name.startswith("Other ") and name.endswith("s"):
                continue          # residual bucket, not a real player
            if pos not in POSITIONS:
                continue
            players[pos].append({
                "Player": name, "Team": team,
                **{out_col: cell.get(wb2.C[src_col]) or 0
                   for src_col, out_col in POSITIONS[pos]["stats"]},
            })
    wb.close()
    return players


# OUTPUT column name -> scoring.py's expected stat-dict key
STAT_KEY = {
    "Pass Att": "pass_att", "Pass Comp": "pass_cmp", "Pass Yds": "pass_yds",
    "Pass TD": "pass_td", "Pass Int": "pass_int",
    "Rush Att": "rush_att", "Rush Yds": "rush_yds", "Rush TD": "rush_td",
    "Targets": "targets", "Rec": "rec", "Rec Yds": "rec_yds",
    "Rec TD": "rec_td",
}


def score(pos, row):
    """One row's stats -> scoring.py's stat-dict -> every scoring format."""
    stats = {STAT_KEY[out_col]: row[out_col]
            for _, out_col in POSITIONS[pos]["stats"]}
    if pos == "QB":
        return {"Fantasy Points": round(scoring.qb_points(stats), 1)}
    out = {
        "Fantasy Points (Half-PPR)" if pos != "WR" else "Fantasy Points (Half)":
            round(scoring.half_ppr_points(stats), 1),
        "Fantasy Points (PPR)": round(scoring.ppr_points(stats), 1),
        "Fantasy Points (STD)": round(scoring.std_points(stats), 1),
    }
    if pos == "TE":
        out["TE Premium"] = round(scoring.te_premium_points(stats), 1)
    return out


def build_tab(ws, pos, rows):
    cols = POSITIONS[pos]["out"]
    widths = {"Player": 24, "Team": 6}
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    for i, name in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 12)
        c = ws.cell(1, i, name)
        c.font = Font(bold=True, size=9, color="FFFFFF")
        c.fill = HDR
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = BOX

    sort_col = POSITIONS[pos]["sort"]
    rows = sorted(rows, key=lambda r: -r.get(sort_col, 0))
    for r, row in enumerate(rows, start=2):
        for i, name in enumerate(cols, start=1):
            v = row.get(name, "")
            c = ws.cell(r, i, v)
            c.border = BOX
            if "Fantasy Points" in name or name == "TE Premium":
                c.number_format = NUM1
                if name == sort_col:
                    c.font = BOLD
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows)+1}"


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--in", dest="src", default="my_projections_v2.xlsx")
    ap.add_argument("--out", default="my_projections_v2_rankings.xlsx")
    ap.add_argument("--min-points", type=float, default=1.0,
                    help="drop players below this many half-PPR/QB points "
                         "(cuts out unseeded zeros)")
    args = ap.parse_args()

    src = os.path.join(HERE, args.src)
    players = read_players(src)

    wb = Workbook()
    wb.remove(wb.active)
    counts = {}
    for pos in ("QB", "RB", "WR", "TE"):
        rows = []
        for p in players.get(pos, []):
            pts = score(pos, p)
            merged = {**p, **pts}
            primary = pts.get(POSITIONS[pos]["sort"], pts.get("Fantasy Points", 0))
            if primary < args.min_points:
                continue
            rows.append(merged)
        build_tab(wb.create_sheet(pos), pos, rows)
        counts[pos] = len(rows)

    out = os.path.join(HERE, args.out)
    wb.save(out)

    print(f"Read {os.path.basename(src)} -> wrote {args.out}")
    for pos, n in counts.items():
        print(f"  {pos:<3} {n:>3} players")
    print(f"\nStatic values, not formulas - re-run this any time you want "
          f"the rankings refreshed. Works even while {args.src} is open in Excel.")


if __name__ == "__main__":
    sys.exit(main())
