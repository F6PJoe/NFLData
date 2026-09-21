#!/usr/bin/env python3
"""Read my_projections_v3.xlsx -> ranked QB/RB/WR/TE tabs, all scoring formats.

Why this recomputes rather than reads
-------------------------------------
In v3 almost everything is a formula - you set shares and rates, the workbook
derives attempts, yards and TDs. openpyxl writes formulas but no cached
results, so a freshly built file has `None` in every calculated cell until
Excel has opened and saved it.

So this script replays the same formula chain in Python from the INPUT cells
(which are real values). Two consequences worth knowing:

  * it works on a brand-new workbook you haven't opened yet
  * it reads the file read-only, so it runs while the workbook is open in Excel

The arithmetic is kept deliberately identical to build_workbook_v3's formulas.
If you change one, change the other - a divergence here would be silent, since
both would still produce plausible-looking numbers.

Scoring comes from ff_draft_proj/scoring.py, imported not reimplemented, so
there is exactly one scoring definition in the repo.

Usage:
    python build_rankings_v3.py
    python build_rankings_v3.py --min-points 20
"""

import argparse
import os
import sys
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                os.pardir, "ff_draft_proj"))
import scoring  # noqa: E402

import build_workbook_v3 as w3
from teams import TEAMS

HERE = os.path.dirname(os.path.abspath(__file__))

HDR = PatternFill("solid", fgColor="2F4F6F")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM0, NUM1, NUM2, PCT = "0", "0.0", "0.00", "0.0%"

TARGET_RATE = 0.954   # matches Settings!B4 in the workbook


def num(cell):
    v = cell.value
    return float(v) if isinstance(v, (int, float)) else 0.0


def read_team(ws, hdr):
    """Replay the team block: plays x pass rate -> attempts; TDs x rate -> split."""
    def t1(label):
        return num(ws.cell(w3.R_T1_VAL, w3.T1[label]))

    def t2(label):
        return num(ws.cell(w3.R_T2_VAL, w3.T2[label]))

    plays = t1("Team Plays")
    pass_rate = t1("Pass Play %")
    pass_att = round(plays * pass_rate)
    total_td = t2("Total TD")
    rush_td_rate = t2("Rush TD %")
    rush_td = total_td * rush_td_rate
    return {
        "pass_att": pass_att,
        "rush_att": plays - pass_att,
        "targets": round(pass_att * TARGET_RATE),
        "rush_td": rush_td,
        "pass_td": total_td - rush_td,
    }


def read_players(path):
    """Every player, with the v3 chain recomputed from his input cells."""
    wb = load_workbook(path, data_only=True, read_only=True)
    out = defaultdict(list)
    for team in TEAMS:
        if team not in wb.sheetnames:
            continue
        ws = wb[team]
        hdr = {c.value: c.column for c in ws[w3.R_HDR] if c.value}
        if "Player" not in hdr:
            continue
        tt = read_team(ws, hdr)

        rows = []
        for r in ws.iter_rows(min_row=w3.R_FIRST, max_row=w3.R_LAST):
            cell = {c.column: c for c in r}

            def g(name):
                col = hdr.get(name)
                return num(cell[col]) if col in cell else 0.0

            name = cell[hdr["Player"]].value if hdr["Player"] in cell else None
            pos = cell[hdr["Pos"]].value if hdr["Pos"] in cell else None
            if not name or not isinstance(name, str):
                continue
            if name.startswith("Other ") and name.endswith("s"):
                continue
            if pos not in ("QB", "RB", "WR", "TE", "FB"):
                continue

            targets = g("Tgt %") * tt["targets"]
            rec = targets * g("Catch %")
            rec_yds = rec * g("Y/R")
            rush_att = g("Rush %") * tt["rush_att"]
            pass_att = g("Pass %") * tt["pass_att"]
            rows.append({
                "player": name, "team": team,
                "pos": "RB" if pos == "FB" else pos,
                "pass_att": pass_att,
                "pass_cmp": pass_att * g("Comp %"),
                "pass_share": g("Pass %"),
                "pass_td": g("Pass %") * tt["pass_td"],
                "pass_int": pass_att * g("Int Rate"),
                "rush_att": rush_att,
                "rush_yds": rush_att * g("Y/C"),
                "rush_td": g("Rush TD %") * tt["rush_td"],
                "targets": targets, "rec": rec, "rec_yds": rec_yds,
                "rec_td": g("Rec TD %") * tt["pass_td"],
            })

        # Passing yards are bottom-up: a QB gets his share of what this team's
        # receivers actually catch. Needs every receiver's yards first, so it
        # happens after the row loop - same dependency the workbook formula has.
        team_rec_yds = sum(p["rec_yds"] for p in rows)
        for p in rows:
            p["pass_yds"] = p["pass_share"] * team_rec_yds
            out[p["pos"]].append(p)

    wb.close()
    return out


def score_row(p):
    stats = {k: p[k] for k in ("pass_yds", "pass_td", "pass_int", "rush_yds",
                               "rush_td", "rec_yds", "rec_td", "rec")}
    return {
        "STD": scoring.qb_points(stats) if p["pos"] == "QB"
        else scoring.std_points(stats),
        "Half": scoring.qb_points(stats) if p["pos"] == "QB"
        else scoring.half_ppr_points(stats),
        "PPR": scoring.qb_points(stats) if p["pos"] == "QB"
        else scoring.ppr_points(stats),
        "TE Premium": (scoring.te_premium_points(stats)
                       if p["pos"] == "TE" else None),
    }


# (column, source key, number format). Rate columns are included so a line can
# be sanity-checked without opening the team tab - a 6.5 Y/C jumps out here.
LAYOUT = {
    "QB": [("Player", "player", None), ("Team", "team", None),
           ("Pass Att", "pass_att", NUM0), ("Cmp", "pass_cmp", NUM0),
           ("Cmp%", "cmp_pct", PCT), ("Pass Yds", "pass_yds", NUM0),
           ("Y/A", "ypa", NUM2), ("Pass TD", "pass_td", NUM1),
           ("Int", "pass_int", NUM1),
           ("Rush Att", "rush_att", NUM0), ("Rush Yds", "rush_yds", NUM0),
           ("Y/C", "ypc", NUM2), ("Rush TD", "rush_td", NUM1),
           ("STD", "STD", NUM1), ("Half", "Half", NUM1), ("PPR", "PPR", NUM1)],
    "SKILL": [("Player", "player", None), ("Team", "team", None),
              ("Targets", "targets", NUM0), ("Rec", "rec", NUM0),
              ("Catch%", "catch", PCT), ("Rec Yds", "rec_yds", NUM0),
              ("Y/R", "ypr", NUM2), ("Rec TD", "rec_td", NUM1),
              ("Rush Att", "rush_att", NUM0), ("Rush Yds", "rush_yds", NUM0),
              ("Y/C", "ypc", NUM2), ("Rush TD", "rush_td", NUM1),
              ("STD", "STD", NUM1), ("Half", "Half", NUM1),
              ("PPR", "PPR", NUM1)],
}


def build_tab(ws, pos, rows):
    cols = list(LAYOUT["QB" if pos == "QB" else "SKILL"])
    if pos == "TE":
        cols.append(("TE Prem", "TE Premium", NUM1))
    cols.insert(0, ("#", "rank", NUM0))

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    widths = {"Player": 23, "Team": 6, "#": 4}
    for i, (name, _, _) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 9)
        c = ws.cell(1, i, name)
        c.font = Font(bold=True, size=9, color="FFFFFF")
        c.fill = HDR
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = BOX

    for r, row in enumerate(rows, start=2):
        for i, (name, key, fmt) in enumerate(cols, start=1):
            c = ws.cell(r, i, row.get(key))
            c.border = BOX
            if fmt:
                c.number_format = fmt
            if name == "Half":
                c.font = BOLD
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows)+1}"


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--in", dest="src", default="my_projections_v3.xlsx")
    ap.add_argument("--out", default="rankings_v3.xlsx")
    ap.add_argument("--min-points", type=float, default=1.0)
    args = ap.parse_args()

    src = os.path.join(HERE, args.src)
    if not os.path.exists(src):
        sys.exit(f"Missing {args.src}. Run build_workbook_v3.py first.")

    players = read_players(src)

    wb = Workbook()
    wb.remove(wb.active)
    counts = {}
    for pos in ("QB", "RB", "WR", "TE"):
        rows = []
        for p in players.get(pos, []):
            pts = score_row(p)
            merged = {**p, **pts}
            # Rates, recomputed from the projected line rather than the input,
            # so they reflect what the projection actually implies.
            merged["ypc"] = (p["rush_yds"] / p["rush_att"]) if p["rush_att"] else None
            merged["ypr"] = (p["rec_yds"] / p["rec"]) if p["rec"] else None
            merged["ypa"] = (p["pass_yds"] / p["pass_att"]) if p["pass_att"] else None
            merged["cmp_pct"] = (p["pass_cmp"] / p["pass_att"]) if p["pass_att"] else None
            merged["catch"] = (p["rec"] / p["targets"]) if p["targets"] else None
            if merged["Half"] < args.min_points:
                continue
            rows.append(merged)
        rows.sort(key=lambda x: -x["Half"])
        for n, row in enumerate(rows, start=1):
            row["rank"] = n
        build_tab(wb.create_sheet(pos), pos, rows)
        counts[pos] = len(rows)

    out = os.path.join(HERE, args.out)
    try:
        wb.save(out)
    except PermissionError:
        sys.exit(f"\nCan't write {args.out} - it's open in Excel.\n")

    print(f"Read {os.path.basename(src)} -> wrote {args.out}")
    for pos, n in counts.items():
        print(f"  {pos:<3} {n:>3} players")
    print("\n  Static values - re-run any time. Works while the workbook is "
          "open in Excel,\n  and on a fresh build you haven't opened yet.")


if __name__ == "__main__":
    sys.exit(main())
